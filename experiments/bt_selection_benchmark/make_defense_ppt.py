#!/usr/bin/env python3
"""Create a formal defense PPT for the BT selection graduation project.

The deck is generated with native PowerPoint shapes, tables and chart-like
objects where practical, so most visual elements remain editable/vector based.
Use --template to start from a university template if available.
"""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
FINAL = ROOT / "experiments" / "bt_selection_benchmark" / "results" / "final_visualizations"
DEFAULT_OUT = FINAL / "bt_selector_defense_presentation.pptx"


BLUE = RGBColor(37, 99, 235)
NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(248, 250, 252)
LINE = RGBColor(203, 213, 225)
GREEN = RGBColor(22, 163, 74)
RED = RGBColor(220, 38, 38)
AMBER = RGBColor(217, 119, 6)
PURPLE = RGBColor(124, 58, 237)
GRAY = RGBColor(100, 116, 139)


def read_csv(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def num(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(out) else out


def first(row: Dict[str, Any], names: Iterable[str], default: str = "") -> str:
    for name in names:
        if name in row and row[name] not in ("", None):
            return str(row[name])
    return default


def short(name: str) -> str:
    return (
        name.replace("transformer_", "trans_")
        .replace("v5_learned_", "learned_")
        .replace("v5_no_", "no_")
        .replace("_", " ")
    )


def set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_line(shape, color: RGBColor = LINE, width: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def text_frame(shape, text: str, size: int = 18, color: RGBColor = NAVY, bold: bool = False, align=None) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Aptos"
    if align is not None:
        p.alignment = align


def add_title(slide, title: str, subtitle: str = "") -> None:
    box = slide.shapes.add_textbox(Cm(1.0), Cm(0.45), Cm(31.8), Cm(1.0))
    text_frame(box, title, 26, NAVY, True)
    if subtitle:
        sub = slide.shapes.add_textbox(Cm(1.05), Cm(1.35), Cm(30.5), Cm(0.55))
        text_frame(sub, subtitle, 11, SLATE)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.0), Cm(1.95), Cm(31.8), Cm(0.04))
    set_fill(line, BLUE)
    line.line.fill.background()


def add_footer(slide, page: int) -> None:
    box = slide.shapes.add_textbox(Cm(1.0), Cm(18.2), Cm(28.5), Cm(0.4))
    text_frame(box, "KIOS + Transformer BT Selector", 8, GRAY)
    num_box = slide.shapes.add_textbox(Cm(31.2), Cm(18.2), Cm(1.4), Cm(0.4))
    text_frame(num_box, str(page), 8, GRAY, align=PP_ALIGN.RIGHT)


def add_bullets(slide, items: Sequence[str], x: float, y: float, w: float, h: float, size: int = 18) -> None:
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = NAVY
        p.font.name = "Aptos"
        p.space_after = Pt(7)


def add_chip(slide, text: str, x: float, y: float, w: float, color: RGBColor = BLUE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(0.68))
    set_fill(shape, color)
    shape.line.fill.background()
    text_frame(shape, text, 10, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)


def add_step(slide, text: str, x: float, y: float, w: float, h: float, color: RGBColor = LIGHT) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    set_fill(shape, color)
    set_line(shape)
    text_frame(shape, text, 13, NAVY, True, PP_ALIGN.CENTER)


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float) -> None:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    conn.line.color.rgb = BLUE
    conn.line.width = Pt(1.6)
    conn.line.end_arrowhead = True


def add_simple_table(slide, headers: Sequence[str], rows: Sequence[Sequence[str]], x: float, y: float, w: float, h: float) -> None:
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Cm(x), Cm(y), Cm(w), Cm(h))
    table = table_shape.table
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        set_fill(cell, BLUE)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            paragraph.font.bold = True
            paragraph.font.size = Pt(10)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            set_fill(cell, RGBColor(255, 255, 255) if r % 2 else LIGHT)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.color.rgb = NAVY
                paragraph.font.size = Pt(9)


def add_bar_chart(slide, rows: List[Tuple[str, float]], x: float, y: float, w: float, h: float, title: str, color: RGBColor) -> None:
    title_box = slide.shapes.add_textbox(Cm(x), Cm(y - 0.65), Cm(w), Cm(0.45))
    text_frame(title_box, title, 14, NAVY, True)
    if not rows:
        add_step(slide, "No data available", x, y, w, h)
        return
    max_value = max(value for _, value in rows) or 1.0
    bar_gap = 0.12
    bar_h = max(0.26, (h - 0.6) / len(rows) - bar_gap)
    label_w = min(4.8, w * 0.32)
    chart_w = w - label_w - 0.9
    for idx, (label, value) in enumerate(rows):
        yy = y + idx * (bar_h + bar_gap)
        lab = slide.shapes.add_textbox(Cm(x), Cm(yy), Cm(label_w), Cm(bar_h))
        text_frame(lab, short(label), 8, NAVY)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w + 0.25), Cm(yy + 0.03), Cm(chart_w), Cm(bar_h * 0.72))
        set_fill(bg, RGBColor(226, 232, 240))
        bg.line.fill.background()
        bw = chart_w * (value / max_value)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w + 0.25), Cm(yy + 0.03), Cm(max(0.03, bw)), Cm(bar_h * 0.72))
        set_fill(bar, color)
        bar.line.fill.background()
        val_box = slide.shapes.add_textbox(Cm(x + label_w + chart_w + 0.35), Cm(yy - 0.02), Cm(1.0), Cm(bar_h))
        text_frame(val_box, f"{value:.2f}", 8, SLATE)


def add_heatmap(slide, matrix_rows: List[Tuple[str, List[Optional[float]]]], cols: Sequence[str], x: float, y: float, w: float, h: float, title: str, reverse: bool = False) -> None:
    title_box = slide.shapes.add_textbox(Cm(x), Cm(y - 0.65), Cm(w), Cm(0.45))
    text_frame(title_box, title, 14, NAVY, True)
    values = [v for _, vals in matrix_rows for v in vals if v is not None]
    if not values:
        add_step(slide, "No data available", x, y, w, h)
        return
    mn, mx = min(values), max(values)
    label_w = 4.4
    cell_w = (w - label_w) / max(1, len(cols))
    cell_h = h / max(1, len(matrix_rows) + 1)
    for c, col in enumerate(cols):
        box = slide.shapes.add_textbox(Cm(x + label_w + c * cell_w), Cm(y), Cm(cell_w), Cm(cell_h))
        text_frame(box, col.replace(" ", "\n"), 7, SLATE, True, PP_ALIGN.CENTER)
    for r, (name, vals) in enumerate(matrix_rows, start=1):
        lab = slide.shapes.add_textbox(Cm(x), Cm(y + r * cell_h), Cm(label_w), Cm(cell_h))
        text_frame(lab, short(name), 8, NAVY)
        for c, value in enumerate(vals):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w + c * cell_w), Cm(y + r * cell_h), Cm(cell_w), Cm(cell_h))
            if value is None:
                set_fill(cell, RGBColor(241, 245, 249))
                set_line(cell, RGBColor(226, 232, 240), 0.5)
                continue
            t = 0.5 if mx == mn else (value - mn) / (mx - mn)
            if reverse:
                t = 1 - t
            color = RGBColor(int(239 - 190 * t), int(246 - 120 * t), int(255 - 60 * t))
            set_fill(cell, color)
            set_line(cell, RGBColor(255, 255, 255), 0.5)
            txt = slide.shapes.add_textbox(Cm(x + label_w + c * cell_w), Cm(y + r * cell_h + 0.05), Cm(cell_w), Cm(cell_h * 0.85))
            text_frame(txt, f"{value:.2f}", 7, NAVY, True, PP_ALIGN.CENTER)


def add_bubble_tradeoff(slide, rows: List[Dict[str, Any]], x: float, y: float, w: float, h: float) -> None:
    title_box = slide.shapes.add_textbox(Cm(x), Cm(y - 0.65), Cm(w), Cm(0.45))
    text_frame(title_box, "Success-Regret Trade-off", 14, NAVY, True)
    data = [(r["selector"], num(r.get("success")), num(r.get("mean_regret") or r.get("regret")), num(r.get("groups"))) for r in rows]
    data = [(s, sy, rg, g or 30) for s, sy, rg, g in data if sy is not None and rg is not None]
    if not data:
        add_step(slide, "No data available", x, y, w, h)
        return
    max_regret = max(rg for _, _, rg, _ in data) or 1
    ax = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    set_fill(ax, RGBColor(255, 255, 255))
    set_line(ax)
    for selector, success, regret, groups in data[:18]:
        px = x + 0.65 + (regret / max_regret) * (w - 1.4)
        py = y + 0.35 + (1.0 - min(1.0, success)) * (h - 0.9)
        radius = max(0.18, min(0.42, 0.12 + groups / 350))
        bubble = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(px - radius / 2), Cm(py - radius / 2), Cm(radius), Cm(radius))
        fill = GRAY if "oracle" in selector else BLUE if "v5" in selector else GREEN if "transformer" in selector or "feature" in selector else AMBER
        set_fill(bubble, fill)
        bubble.line.fill.background()
        lab = slide.shapes.add_textbox(Cm(px + 0.1), Cm(py - 0.15), Cm(3.0), Cm(0.32))
        text_frame(lab, short(selector), 7, SLATE)
    xlab = slide.shapes.add_textbox(Cm(x + w / 2 - 2), Cm(y + h + 0.1), Cm(4.0), Cm(0.4))
    text_frame(xlab, "Mean regret (lower is better)", 8, SLATE, align=PP_ALIGN.CENTER)
    ylab = slide.shapes.add_textbox(Cm(x - 0.2), Cm(y - 0.45), Cm(4.0), Cm(0.35))
    text_frame(ylab, "Success rate", 8, SLATE)


def load_selector_comparison() -> List[Dict[str, str]]:
    table = FINAL / "final_selector_comparison.csv"
    rows = read_csv(table)
    if rows:
        return rows
    return []


def selector_matrix_rows(value: str) -> Tuple[List[str], List[Tuple[str, List[Optional[float]]]]]:
    rows = []
    for source_file in [
        ("V1", ROOT / "experiments/bt_selection_benchmark/results/summary_by_selector_v1_supported_by_with_sim.csv"),
        ("V2", ROOT / "experiments/bt_selection_benchmark/results/summary_by_selector_v2_stacking_probe_with_sim.csv"),
        ("V2B", ROOT / "experiments/bt_selection_benchmark/results/summary_by_selector_v2b_physical_strategy_with_sim.csv"),
        ("V6", ROOT / "experiments/bt_selection_benchmark/results/summary_by_selector_v6_hard_cases_with_sim.csv"),
        ("V4B", ROOT / "experiments/bt_selection_benchmark/results/v4b_transformer_selector_summary.csv"),
        ("V5", ROOT / "experiments/bt_selection_benchmark/results/v5_ensemble_selector_summary.csv"),
    ]:
        source, path = source_file
        for row in read_csv(path):
            selector = first(row, ["selector"])
            metric = first(row, [value, f"mean_{value}", "mean_regret" if value == "regret" else "success"])
            rows.append((source, selector, num(metric)))
    cols = list(dict.fromkeys(r[0] for r in rows))
    selectors_priority = ["feature_only", "transformer_only", "transformer_fused", "symbolic_only", "shortest_tree", "rule_based", "v5_ensemble", "oracle"]
    all_selectors = list(dict.fromkeys(r[1] for r in rows if r[2] is not None))
    selectors = [s for s in selectors_priority if s in all_selectors]
    selectors.extend([s for s in all_selectors if s not in selectors][: max(0, 10 - len(selectors))])
    lookup = {(src, sel): val for src, sel, val in rows}
    return cols, [(sel, [lookup.get((col, sel)) for col in cols]) for sel in selectors]


def create_deck(template: Optional[Path] = None) -> Presentation:
    prs = Presentation(str(template)) if template and template.exists() else Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]
    return prs


def blank_layout(prs: Presentation):
    return prs.slide_layouts[6]


def slide_title(prs: Presentation, page: int, args) -> int:
    slide = prs.slides.add_slide(blank_layout(prs))
    title = slide.shapes.add_textbox(Cm(1.4), Cm(3.0), Cm(30.5), Cm(2.1))
    text_frame(title, "基于 KIOS 与 Transformer 编码的\n机器人行为树候选选择方法", 31, NAVY, True)
    subtitle = slide.shapes.add_textbox(Cm(1.5), Cm(5.55), Cm(24), Cm(0.8))
    text_frame(subtitle, "LLM-generated Behavior Tree Selection for Robot Task Planning", 16, BLUE, True)
    add_chip(slide, "KIOS", 1.5, 7.1, 2.5, BLUE)
    add_chip(slide, "Behavior Trees", 4.3, 7.1, 3.9, GREEN)
    add_chip(slide, "Transformer Encoder", 8.5, 7.1, 5.0, PURPLE)
    add_chip(slide, "Isaac Gym", 13.8, 7.1, 3.4, AMBER)
    info = slide.shapes.add_textbox(Cm(1.5), Cm(14.8), Cm(15.0), Cm(1.4))
    text_frame(info, f"{args.author}\n{args.department}\n{args.date}", 13, SLATE)
    add_footer(slide, page)
    return page + 1


def build_deck(args) -> Presentation:
    prs = create_deck(args.template)
    page = 1
    selector_rows = load_selector_comparison()
    coverage = read_csv(FINAL / "enhanced_candidate_simulation_outcomes.csv") or read_csv(FINAL / "final_dataset_coverage.csv")
    ablation = read_csv(FINAL / "final_v5_learned_ablation.csv")

    page = slide_title(prs, page, args)

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "研究背景", "从语言任务到机器人可执行行为")
    add_bullets(slide, [
        "LLM 可以根据自然语言生成任务计划，但输出并不天然可执行。",
        "行为树适合表达机器人任务的层次化执行逻辑。",
        "生成的 BT 可能存在结构错误、符号条件缺失或物理执行不稳定。",
        "因此需要从多个候选 BT 中选择更可靠的执行方案。"], 1.3, 3.0, 14.2, 8.5, 18)
    add_step(slide, "Task Instruction", 18.0, 3.2, 4.8, 1.3); add_step(slide, "Behavior Tree", 18.0, 6.0, 4.8, 1.3); add_step(slide, "Robot Execution", 18.0, 8.8, 4.8, 1.3)
    add_arrow(slide, 20.4, 4.5, 20.4, 6.0); add_arrow(slide, 20.4, 7.3, 20.4, 8.8)
    add_step(slide, "Reliability gap", 24.2, 6.0, 5.8, 1.3, RGBColor(254, 243, 199))
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "问题定义", "Multi-candidate behavior tree selection")
    add_step(slide, "Task I", 2.0, 4.0, 3.6, 1.2); add_step(slide, "State S0", 2.0, 6.0, 3.6, 1.2)
    add_step(slide, "Candidates C = {c1, c2, ..., cn}", 8.0, 4.9, 8.2, 1.3)
    add_step(slide, "score(ci)", 18.5, 4.9, 4.2, 1.3)
    add_step(slide, "c* = argmax score(ci)", 25.0, 4.9, 5.8, 1.3, RGBColor(220, 252, 231))
    add_arrow(slide, 5.6, 4.6, 8.0, 5.4); add_arrow(slide, 5.6, 6.6, 8.0, 5.7); add_arrow(slide, 16.2, 5.55, 18.5, 5.55); add_arrow(slide, 22.7, 5.55, 25.0, 5.55)
    add_bullets(slide, ["目标不是判断单个 BT 是否可行，而是在同一任务组内选择相对最优候选。", "关键挑战：符号成功不等于物理稳定，BT 结构合理不等于执行可靠。"], 2.0, 10.5, 27.5, 3.2, 18)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "系统总体架构", "Candidate generation, symbolic evaluation, learned selection and simulation validation")
    steps = [("Task + State", 1.2), ("BT Candidates", 6.0), ("KIOS Evaluation", 11.1), ("Transformer Encoder", 16.6), ("V5 Selector", 22.2), ("Selected BT", 27.0)]
    for text, x in steps:
        add_step(slide, text, x, 5.0, 4.2, 1.35)
    for i in range(len(steps) - 1):
        add_arrow(slide, steps[i][1] + 4.2, 5.68, steps[i + 1][1], 5.68)
    add_step(slide, "Isaac Gym Validation", 22.2, 9.0, 5.4, 1.35, RGBColor(239, 246, 255))
    add_arrow(slide, 29.0, 6.35, 24.9, 9.0)
    add_bullets(slide, ["核心思想：不依赖单个生成结果，而是通过多源评价信号选择候选。", "KIOS 提供低成本符号反馈，Isaac Gym 提供物理执行验证。"], 2.0, 12.7, 27.0, 2.2, 17)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "相关工作与任务启发", "From LLM-generated BTs to relational long-horizon planning")
    add_simple_table(slide, ["Reference", "Role in this project"], [
        ["LLM-as-BT-Planner", "LLM 到行为树生成流程"],
        ["Points2Plans", "关系状态与长时序任务规划启发"],
        ["MuST / Transformer", "技能序列与 token 表示学习启发"],
        ["本文方法", "多候选 BT 生成 + KIOS + Transformer + V5 selector"],
    ], 2.0, 3.1, 28.6, 7.0)
    add_bullets(slide, ["本项目不是直接复现单篇论文，而是将相关思想组合到 BT candidate selection 问题中。"], 2.2, 11.6, 27.0, 1.4, 18)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "KIOS 符号执行", "Low-cost symbolic reliability signals")
    add_simple_table(slide, ["Feature", "Meaning"], [
        ["symbolic_success", "符号执行是否成功"],
        ["goal_satisfaction", "目标满足程度"],
        ["precondition_coverage", "前置条件覆盖率"],
        ["invalid_action_count", "无效动作数量"],
        ["condition_failure_count", "条件失败次数"],
        ["tree_size / tree_depth", "BT 结构复杂度"],
    ], 1.6, 3.0, 15.0, 8.5)
    add_step(slide, "KIOS can filter clearly invalid BTs,\nbut cannot fully model physical stability.", 19.0, 5.2, 10.4, 2.3, RGBColor(239, 246, 255))
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "BT Token Transformer Encoder", "Encoding behavior tree structure")
    chain = [("BT JSON", 2.0), ("Preorder Tokens", 7.2), ("Token IDs", 13.2), ("Transformer Encoder", 18.2), ("BT Embedding", 25.4)]
    for text, x in chain:
        add_step(slide, text, x, 5.0, 4.6 if x != 18.2 else 5.4, 1.35)
    for i in range(len(chain) - 1):
        add_arrow(slide, chain[i][1] + (5.4 if chain[i][1] == 18.2 else 4.6), 5.7, chain[i + 1][1], 5.7)
    add_bullets(slide, ["Transformer-only: uses only BT token structure.", "Transformer-fused: combines BT embedding with task/world/KIOS features.", "定位：结构表示模块，而不是唯一决策模块。"], 3.0, 10.2, 26.0, 4.0, 18)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "V5 Hybrid Learned Selector", "Multi-source reliability fusion")
    add_step(slide, "Selector Agreement", 2.4, 4.5, 6.0, 1.3, RGBColor(239, 246, 255))
    add_step(slide, "Symbolic Reliability", 2.4, 7.0, 6.0, 1.3, RGBColor(240, 253, 244))
    add_step(slide, "Physical Reliability", 2.4, 9.5, 6.0, 1.3, RGBColor(255, 247, 237))
    add_step(slide, "Learned V5 Fusion", 13.0, 7.0, 6.2, 1.45, RGBColor(245, 243, 255))
    add_step(slide, "Final Candidate Score", 24.0, 7.0, 6.0, 1.45, RGBColor(220, 252, 231))
    add_arrow(slide, 8.4, 5.15, 13.0, 7.35); add_arrow(slide, 8.4, 7.65, 13.0, 7.65); add_arrow(slide, 8.4, 10.15, 13.0, 8.0); add_arrow(slide, 19.2, 7.72, 24.0, 7.72)
    eq = slide.shapes.add_textbox(Cm(7.0), Cm(12.2), Cm(20.0), Cm(1.0))
    text_frame(eq, "scoreV5(c) = w1 Rvote(c) + w2 Rsym(c) + w3 Rphy(c)", 20, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "Benchmark 设计", "Progressively harder BT selection tasks")
    add_simple_table(slide, ["Benchmark", "Goal", "Why it matters"], [
        ["V1 supported-by", "基础支持关系", "验证基本放置任务"],
        ["V2 stacking probe", "堆叠任务", "测试长动作链"],
        ["V2B physical strategy", "物理策略差异", "区分中心/边缘等放置策略"],
        ["V6 hard cases", "困难物理约束", "验证复杂场景鲁棒性"],
    ], 1.5, 3.0, 17.0, 6.8)
    outcome_rows = []
    for r in coverage:
        success = num(first(r, ["sim_success", "sim_done"], "0")) or 0
        failure = num(first(r, ["sim_failure"], "0")) or 0
        outcome_rows.append((first(r, ["benchmark"], ""), success + failure))
    add_bar_chart(slide, outcome_rows[:6], 20.0, 4.0, 11.5, 6.3, "Evaluated simulation rows", BLUE)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "总体结果", "Success and regret trade-off")
    add_bubble_tradeoff(slide, selector_rows, 1.5, 3.0, 18.0, 10.5)
    add_bullets(slide, ["Mean regret 越低越接近 oracle。", "V5 方法相比简单启发式更接近 oracle。", "shortest_tree 表现较弱，说明 BT 越短不代表越可靠。"], 21.0, 4.2, 10.2, 6.2, 17)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "不同任务难度下的表现", "Selector regret heatmap")
    cols, rows = selector_matrix_rows("regret")
    add_heatmap(slide, rows, cols, 1.5, 3.0, 22.0, 9.0, "Mean regret by benchmark", reverse=True)
    add_bullets(slide, ["Heatmap 用于观察 selector 在不同 benchmark 上的稳定性。", "困难任务中，融合物理可靠性的信息更重要。"], 24.5, 5.0, 7.2, 4.0, 16)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "V5 消融实验", "Contribution of each reliability source")
    abl_rows = []
    for r in ablation[:10]:
        regret = num(first(r, ["mean_regret", "regret"], ""))
        if regret is not None:
            abl_rows.append((first(r, ["selector"], ""), regret))
    add_bar_chart(slide, abl_rows, 1.5, 3.2, 16.0, 9.8, "Mean regret after ablation", PURPLE)
    add_bullets(slide, ["去除 symbolic reliability 后，符号不可执行候选更难被抑制。", "去除 simulation reliability 后，物理策略差异更难被区分。", "完整 V5 的优势来自多源信号融合。"], 19.0, 4.0, 12.0, 5.8, 17)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "系统 Demo", "End-to-end BT selection demonstration")
    add_step(slide, "Task definition", 2.0, 4.0, 5.2, 1.3)
    add_step(slide, "Candidate BTs", 9.0, 4.0, 5.2, 1.3)
    add_step(slide, "Selector scores", 16.0, 4.0, 5.2, 1.3)
    add_step(slide, "Selected BT", 23.0, 4.0, 5.2, 1.3, RGBColor(220, 252, 231))
    add_arrow(slide, 7.2, 4.65, 9.0, 4.65); add_arrow(slide, 14.2, 4.65, 16.0, 4.65); add_arrow(slide, 21.2, 4.65, 23.0, 4.65)
    add_step(slide, "Web page: http://localhost:8091/", 4.0, 8.0, 11.5, 1.3, RGBColor(239, 246, 255))
    add_step(slide, "Top-down placement view", 17.0, 8.0, 6.5, 1.3, RGBColor(255, 247, 237))
    add_step(slide, "Isaac Gym replay command", 17.0, 10.2, 6.5, 1.3, RGBColor(255, 247, 237))
    add_bullets(slide, ["Demo 展示完整链路：任务、多候选 BT、评分、选择和仿真验证。"], 4.0, 13.1, 24.0, 1.6, 18)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "总结与贡献")
    add_bullets(slide, [
        "构建了 LLM/BT/KIOS/Isaac Gym 结合的机器人任务规划系统。",
        "设计了多组 BT selection benchmark，从基础关系到困难物理约束。",
        "实现了 BT token Transformer encoder，用于行为树结构表示。",
        "提出 V5 hybrid learned selector，融合 selector agreement、symbolic reliability 和 physical reliability。",
        "通过仿真和消融实验验证多源融合的有效性。"], 2.0, 3.3, 28.0, 7.8, 18)
    add_footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank_layout(prs)); add_title(slide, "未来工作")
    add_bullets(slide, [
        "补充更大规模 GPT single-vs-multi baseline。",
        "增加更多真实机器人或复杂仿真任务。",
        "尝试更大规模 Transformer 或图结构模型。",
        "引入闭环 BT repair 与在线反馈。",
        "将仿真结果用于主动学习和数据扩充。"], 2.0, 3.4, 28.0, 7.0, 19)
    add_footer(slide, page); page += 1

    return prs


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--template", type=Path, default=None, help="Optional university PPTX template.")
    parser.add_argument("--author", default="Your Name")
    parser.add_argument("--department", default="Department / School")
    parser.add_argument("--date", default="2026")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    prs = build_deck(args)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print({"result": "success", "output": str(args.output), "slides": len(prs.slides)})


if __name__ == "__main__":
    main()
