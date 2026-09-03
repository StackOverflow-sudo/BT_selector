#!/usr/bin/env python3
"""Create an enhanced English defense deck from an existing PPTX/template.

The generated deck uses native PowerPoint text, tables, shapes, and chart-like
elements as much as possible, so the content remains editable and vector based.
"""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
FINAL = ROOT / "experiments" / "bt_selection_benchmark" / "results" / "final_visualizations"
DEFAULT_INPUT = Path("/mnt/c/Users/lenovo/Desktop/bt_selector_defense_presentation.pptx")
DEFAULT_OUTPUT = Path("/mnt/c/Users/lenovo/Desktop/bt_selector_defense_presentation_enhanced_en.pptx")

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(8, 145, 178)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)
PURPLE = RGBColor(124, 58, 237)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(100, 116, 139)
GRAY = RGBColor(100, 116, 139)
LINE = RGBColor(203, 213, 225)
PALE = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)


def read_csv(path: Path) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def n(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(out) else out


def first(row: Dict[str, Any], names: Iterable[str], default: str = "") -> str:
    for name in names:
        if name in row and row[name] not in ("", None):
            return str(row[name])
    return default


def short(name: str) -> str:
    return (
        name.replace("v5_learned_", "Learned ")
        .replace("transformer_", "Transformer ")
        .replace("feature_only", "Feature-only")
        .replace("symbolic_only", "Symbolic-only")
        .replace("shortest_tree", "Shortest-tree")
        .replace("v5_ensemble", "Manual V5")
        .replace("oracle", "Oracle")
        .replace("_", " ")
    )


def rgb_lerp(a: RGBColor, b: RGBColor, t: float) -> RGBColor:
    t = max(0.0, min(1.0, t))
    return RGBColor(
        int(a[0] + (b[0] - a[0]) * t),
        int(a[1] + (b[1] - a[1]) * t),
        int(a[2] + (b[2] - a[2]) * t),
    )


def clear_slides(prs: Presentation) -> None:
    while len(prs.slides) > 0:
        r_id = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]


def new_prs(template: Path) -> Presentation:
    prs = Presentation(str(template)) if template.exists() else Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)
    clear_slides(prs)
    return prs


def blank(prs: Presentation):
    return prs.slide_layouts[6]


def fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color: RGBColor = LINE, width: float = 0.8) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def textbox(slide, text: str, x: float, y: float, w: float, h: float, size: int = 16, color: RGBColor = NAVY, bold: bool = False, align=None):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Cm(0.05)
    tf.margin_right = Cm(0.05)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def title(slide, text: str, subtitle: str = "") -> None:
    textbox(slide, text, 1.0, 0.45, 28.5, 0.85, 25, NAVY, True)
    if subtitle:
        textbox(slide, subtitle, 1.05, 1.27, 28.0, 0.5, 10, MUTED)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.0), Cm(1.88), Cm(31.8), Cm(0.05))
    fill(bar, BLUE)
    bar.line.fill.background()


def footer(slide, page: int) -> None:
    textbox(slide, "KIOS + Transformer BT Selector", 1.0, 18.15, 12.0, 0.35, 8, MUTED)
    textbox(slide, str(page), 31.3, 18.15, 1.2, 0.35, 8, MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, items: Sequence[str], x: float, y: float, w: float, h: float, size: int = 16) -> None:
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = NAVY
        p.space_after = Pt(8)


def card(slide, heading: str, body: str, x: float, y: float, w: float, h: float, color: RGBColor = BLUE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    fill(shape, WHITE)
    line(shape, LINE, 1.0)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(0.12), Cm(h))
    fill(stripe, color)
    stripe.line.fill.background()
    textbox(slide, heading, x + 0.35, y + 0.25, w - 0.6, 0.45, 13, color, True)
    textbox(slide, body, x + 0.35, y + 0.85, w - 0.6, h - 1.0, 10, SLATE)


def step(slide, text: str, x: float, y: float, w: float, h: float, color: RGBColor = PALE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    fill(shape, color)
    line(shape, LINE, 1.0)
    textbox(slide, text, x + 0.15, y + 0.18, w - 0.3, h - 0.25, 11, NAVY, True, PP_ALIGN.CENTER)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = BLUE) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.6)
    c.line.end_arrowhead = True


def table(slide, headers: Sequence[str], rows: Sequence[Sequence[str]], x: float, y: float, w: float, h: float) -> None:
    t = slide.shapes.add_table(len(rows) + 1, len(headers), Cm(x), Cm(y), Cm(w), Cm(h)).table
    for c, head in enumerate(headers):
        cell = t.cell(0, c)
        cell.text = head
        fill(cell, NAVY)
        for p in cell.text_frame.paragraphs:
            p.font.name = "Aptos"
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            fill(cell, WHITE if r % 2 else PALE)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(8)
                p.font.color.rgb = NAVY


def bar_chart(slide, data: List[Tuple[str, float]], x: float, y: float, w: float, h: float, title_text: str, color: RGBColor, lower_is_better: bool = False) -> None:
    textbox(slide, title_text, x, y - 0.55, w, 0.35, 12, NAVY, True)
    if not data:
        step(slide, "No data available", x, y, w, h)
        return
    max_v = max(v for _, v in data) or 1.0
    label_w = 4.9
    bar_area = w - label_w - 0.9
    row_h = h / max(1, len(data))
    for i, (label, value) in enumerate(data):
        yy = y + i * row_h
        textbox(slide, short(label), x, yy + 0.02, label_w, row_h * 0.75, 7, NAVY)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w), Cm(yy + 0.08), Cm(bar_area), Cm(row_h * 0.45))
        fill(bg, RGBColor(226, 232, 240))
        bg.line.fill.background()
        ratio = value / max_v
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w), Cm(yy + 0.08), Cm(max(0.04, bar_area * ratio)), Cm(row_h * 0.45))
        fill(bar, color if not lower_is_better else rgb_lerp(GREEN, RED, ratio))
        bar.line.fill.background()
        textbox(slide, f"{value:.2f}", x + label_w + bar_area + 0.15, yy - 0.01, 0.8, row_h * 0.75, 7, SLATE)


def heatmap(slide, cols: List[str], rows: List[Tuple[str, List[Optional[float]]]], x: float, y: float, w: float, h: float, title_text: str, lower_is_better: bool = False) -> None:
    textbox(slide, title_text, x, y - 0.55, w, 0.35, 12, NAVY, True)
    vals = [v for _, line_vals in rows for v in line_vals if v is not None]
    if not vals:
        step(slide, "No data available", x, y, w, h)
        return
    mn, mx = min(vals), max(vals)
    label_w = 4.2
    cell_w = (w - label_w) / max(1, len(cols))
    cell_h = h / max(1, len(rows) + 1)
    for c, col in enumerate(cols):
        textbox(slide, col, x + label_w + c * cell_w, y, cell_w, cell_h, 7, SLATE, True, PP_ALIGN.CENTER)
    for r, (label, line_vals) in enumerate(rows, start=1):
        textbox(slide, short(label), x, y + r * cell_h + 0.03, label_w, cell_h * 0.8, 7, NAVY)
        for c, value in enumerate(line_vals):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x + label_w + c * cell_w), Cm(y + r * cell_h), Cm(cell_w), Cm(cell_h))
            if value is None:
                fill(cell, RGBColor(241, 245, 249))
            else:
                t = 0.5 if mx == mn else (value - mn) / (mx - mn)
                if lower_is_better:
                    t = 1 - t
                fill(cell, rgb_lerp(RGBColor(254, 226, 226), RGBColor(22, 163, 74), t))
                textbox(slide, f"{value:.2f}", x + label_w + c * cell_w, y + r * cell_h + 0.06, cell_w, cell_h * 0.8, 7, NAVY, True, PP_ALIGN.CENTER)
            line(cell, WHITE, 0.5)


def load_final_selector_rows() -> List[Dict[str, str]]:
    return read_csv(FINAL / "final_selector_comparison.csv")


def load_ablation_rows() -> List[Dict[str, str]]:
    return read_csv(FINAL / "final_v5_learned_ablation.csv")


def load_outcomes() -> List[Dict[str, str]]:
    return read_csv(FINAL / "enhanced_candidate_simulation_outcomes.csv") or read_csv(FINAL / "final_dataset_coverage.csv")


def selector_matrix(metric: str) -> Tuple[List[str], List[Tuple[str, List[Optional[float]]]]]:
    sources = [
        ("V1", "summary_by_selector_v1_supported_by_with_sim.csv"),
        ("V2", "summary_by_selector_v2_stacking_probe_with_sim.csv"),
        ("V2B", "summary_by_selector_v2b_physical_strategy_with_sim.csv"),
        ("V6", "summary_by_selector_v6_hard_cases_with_sim.csv"),
        ("V4B", "v4b_transformer_selector_summary.csv"),
        ("V5", "v5_ensemble_selector_summary.csv"),
    ]
    triples = []
    for source, fname in sources:
        for row in read_csv(ROOT / "experiments" / "bt_selection_benchmark" / "results" / fname):
            selector = first(row, ["selector"])
            value = first(row, [metric, "mean_regret" if metric == "regret" else "success"])
            triples.append((source, selector, n(value)))
    cols = list(dict.fromkeys(t[0] for t in triples))
    priority = ["feature_only", "transformer_only", "transformer_fused", "symbolic_only", "shortest_tree", "rule_based", "v5_ensemble", "oracle"]
    all_selectors = list(dict.fromkeys(t[1] for t in triples if t[2] is not None))
    selectors = [s for s in priority if s in all_selectors]
    selectors.extend([s for s in all_selectors if s not in selectors][: max(0, 10 - len(selectors))])
    lookup = {(src, sel): val for src, sel, val in triples}
    return cols, [(sel, [lookup.get((col, sel)) for col in cols]) for sel in selectors]


def title_slide(prs: Presentation, page: int, args) -> int:
    slide = prs.slides.add_slide(blank(prs))
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0), Cm(0), Cm(33.867), Cm(19.05))
    fill(bg, RGBColor(241, 245, 249))
    bg.line.fill.background()
    textbox(slide, "KIOS + Transformer\nBehavior Tree Selector", 1.3, 2.5, 21.0, 2.4, 32, NAVY, True)
    textbox(slide, "A multi-candidate BT selection system for robot task planning", 1.4, 5.7, 21.0, 0.7, 16, BLUE, True)
    card(slide, "Core Idea", "Generate multiple BTs, evaluate them symbolically, encode their structure, and select the most reliable one.", 1.4, 8.2, 12.5, 3.2, BLUE)
    card(slide, "Final Method", "V5 learned ensemble combining selector agreement, symbolic reliability, and physical reliability.", 15.0, 8.2, 12.5, 3.2, PURPLE)
    textbox(slide, f"{args.author}\n{args.department}\n{args.date}", 1.4, 15.0, 13.0, 1.3, 13, SLATE)
    footer(slide, page)
    return page + 1


def build_deck(args) -> Presentation:
    prs = new_prs(args.input)
    page = 1
    selectors = load_final_selector_rows()
    ablations = load_ablation_rows()
    outcomes = load_outcomes()

    page = title_slide(prs, page, args)

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Motivation", "LLM plans are useful, but robot execution needs reliability")
    card(slide, "Problem", "A single generated behavior tree can be structurally valid but symbolically incomplete or physically unstable.", 1.4, 3.0, 9.2, 4.0, RED)
    card(slide, "Opportunity", "Generating multiple BT candidates gives the system a choice space instead of relying on one fragile output.", 12.0, 3.0, 9.2, 4.0, GREEN)
    card(slide, "Research Question", "How can we select the best BT candidate using symbolic, structural, and physical evidence?", 22.6, 3.0, 9.2, 4.0, BLUE)
    bullets(slide, ["Behavior Trees provide interpretable task execution logic.", "KIOS checks symbolic executability before expensive simulation.", "Isaac Gym reveals physical differences that symbolic execution can miss."], 2.0, 9.0, 28.0, 4.0, 18)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Problem Formulation", "Select one reliable BT from a candidate set")
    step(slide, "Task Instruction I", 1.7, 4.0, 4.5, 1.2); step(slide, "Initial State S0", 1.7, 6.0, 4.5, 1.2)
    step(slide, "Candidate Set\nC = {c1, c2, ..., cn}", 8.2, 4.9, 6.2, 1.5, RGBColor(239, 246, 255))
    step(slide, "Score F(I, S0, ci)", 17.0, 4.9, 5.5, 1.5, RGBColor(245, 243, 255))
    step(slide, "Selected BT\nc* = argmax F", 25.0, 4.9, 5.5, 1.5, RGBColor(220, 252, 231))
    arrow(slide, 6.2, 4.6, 8.2, 5.45); arrow(slide, 6.2, 6.6, 8.2, 5.75); arrow(slide, 14.4, 5.65, 17.0, 5.65); arrow(slide, 22.5, 5.65, 25.0, 5.65)
    bullets(slide, ["The target is not only symbolic success.", "Two BTs can satisfy the same symbolic goal while producing different physical outcomes.", "Mean regret measures how far the selector is from the oracle choice."], 2.0, 10.0, 27.5, 4.0, 18)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "System Architecture", "End-to-end pipeline")
    nodes = [("Task + State", 1.4), ("BT Candidates", 6.5), ("KIOS Metrics", 11.8), ("BT Encoder", 17.1), ("V5 Selector", 22.4), ("Selected BT", 27.3)]
    for label, x in nodes:
        step(slide, label, x, 5.0, 4.2, 1.3)
    for i in range(len(nodes) - 1):
        arrow(slide, nodes[i][1] + 4.2, 5.65, nodes[i + 1][1], 5.65)
    step(slide, "Isaac Gym\nValidation", 21.3, 8.6, 5.4, 1.4, RGBColor(255, 247, 237))
    arrow(slide, 29.0, 6.3, 24.0, 8.6, AMBER)
    bullets(slide, ["Low-cost symbolic screening is performed before simulation-heavy validation.", "The final selector combines model agreement, symbolic reliability, and physical reliability."], 2.0, 12.2, 27.5, 2.4, 18)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Candidate Generation", "From task instruction to multiple BT options")
    table(slide, ["Source", "Use in this project"], [
        ["Rule-based BTs", "Controlled benchmark candidates"],
        ["Existing benchmark BTs", "Repeatable selector evaluation"],
        ["GPT-generated BTs", "Online candidate generation demo"],
    ], 1.6, 3.0, 14.0, 4.8)
    step(slide, "Single BT generation\nhas no fallback.", 19.0, 3.6, 7.0, 1.4, RGBColor(254, 226, 226))
    step(slide, "Multi-candidate generation\ncreates a selection space.", 19.0, 6.2, 7.0, 1.4, RGBColor(220, 252, 231))
    bullets(slide, ["The online GPT baseline was implemented, but API rate limits make it a pilot result rather than a main quantitative claim.", "The main thesis results use completed benchmark and simulation data."], 2.0, 10.2, 27.0, 3.1, 17)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "KIOS Symbolic Evaluation", "Fast reliability signals before simulation")
    table(slide, ["Metric", "Interpretation"], [
        ["symbolic_success", "Does the BT reach the symbolic goal?"],
        ["goal_satisfaction", "Fraction of goal predicates satisfied"],
        ["precondition_coverage", "Whether required checks are covered"],
        ["invalid_action_count", "Unsupported or impossible actions"],
        ["condition_failure_count", "Failed condition checks"],
        ["tree_size / tree_depth", "BT structural complexity"],
    ], 1.5, 3.0, 16.0, 8.2)
    card(slide, "Why KIOS matters", "It is cheaper than simulation and helps reject clearly invalid BTs early.", 19.5, 4.0, 9.5, 2.6, BLUE)
    card(slide, "Limitation", "Symbolic success alone cannot judge placement stability or object displacement.", 19.5, 7.2, 9.5, 2.6, AMBER)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "BT Token Transformer Encoder", "Vector-based structure representation")
    chain = [("BT JSON", 1.8), ("Preorder\nTokens", 7.0), ("Token IDs", 12.3), ("Transformer\nEncoder", 17.4), ("BT\nEmbedding", 24.0)]
    for label, x in chain:
        step(slide, label, x, 5.0, 4.4, 1.4)
    for i in range(len(chain) - 1):
        arrow(slide, chain[i][1] + 4.4, 5.7, chain[i + 1][1], 5.7, PURPLE)
    card(slide, "Transformer-only", "Uses only the BT token sequence.", 3.0, 9.6, 8.0, 2.5, PURPLE)
    card(slide, "Transformer-fused", "Combines BT embedding with task, world, and KIOS features.", 12.8, 9.6, 9.0, 2.5, GREEN)
    card(slide, "Interpretation", "The encoder is a structural feature module, not the only decision-maker.", 23.2, 9.6, 7.5, 2.5, BLUE)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "V5 Hybrid Learned Selector", "Final algorithm")
    step(slide, "Selector\nAgreement", 2.0, 4.0, 5.6, 1.4, RGBColor(239, 246, 255))
    step(slide, "Symbolic\nReliability", 2.0, 6.7, 5.6, 1.4, RGBColor(240, 253, 244))
    step(slide, "Physical\nReliability", 2.0, 9.4, 5.6, 1.4, RGBColor(255, 247, 237))
    step(slide, "Learned Fusion", 13.6, 6.6, 6.0, 1.6, RGBColor(245, 243, 255))
    step(slide, "Final BT Score", 24.4, 6.6, 5.8, 1.6, RGBColor(220, 252, 231))
    arrow(slide, 7.6, 4.7, 13.6, 7.0); arrow(slide, 7.6, 7.4, 13.6, 7.4); arrow(slide, 7.6, 10.1, 13.6, 7.8); arrow(slide, 19.6, 7.4, 24.4, 7.4)
    textbox(slide, "scoreV5(c) = w1 Rvote(c) + w2 Rsym(c) + w3 Rphy(c)", 5.2, 13.2, 23.5, 0.8, 20, NAVY, True, PP_ALIGN.CENTER)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Benchmark Design", "Progressively harder task groups")
    table(slide, ["Benchmark", "Goal", "Purpose"], [
        ["V1 supported-by", "Basic support relation", "Check basic placing behavior"],
        ["V2 stacking probe", "Stacking tasks", "Test longer action chains"],
        ["V2B physical strategy", "Stable vs risky placement", "Separate symbolic and physical quality"],
        ["V6 hard cases", "Hard physical constraints", "Stress-test the final selector"],
    ], 1.4, 3.0, 19.0, 5.7)
    outcome_data = []
    for r in outcomes:
        label = first(r, ["benchmark"])
        val = n(first(r, ["sim_success"])) or n(first(r, ["sim_done"])) or 0
        outcome_data.append((label, val))
    bar_chart(slide, outcome_data[:6], 22.0, 4.0, 9.8, 5.6, "Simulation coverage", BLUE)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Overall Results", "Selector success-regret trade-off")
    plot_rows = []
    for r in selectors:
        success = n(first(r, ["success"]))
        regret = n(first(r, ["mean_regret", "regret"]))
        groups = n(first(r, ["groups"])) or 30
        if success is not None and regret is not None:
            plot_rows.append((first(r, ["selector"]), success, regret, groups))
    # Draw a native scatter-like chart.
    x, y, w, h = 1.6, 3.0, 18.5, 10.4
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    fill(bg, WHITE); line(bg)
    max_reg = max([p[2] for p in plot_rows] or [1])
    for label, success, regret, groups in plot_rows[:18]:
        px = x + 0.6 + regret / max_reg * (w - 1.2)
        py = y + 0.5 + (1 - min(1, success)) * (h - 1.0)
        size = max(0.22, min(0.55, groups / 230))
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(px), Cm(py), Cm(size), Cm(size))
        fill(dot, GRAY if "oracle" in label else PURPLE if "v5" in label else GREEN if "transformer" in label or "feature" in label else AMBER)
        dot.line.fill.background()
        textbox(slide, short(label), px + 0.18, py - 0.05, 3.5, 0.3, 7, SLATE)
    textbox(slide, "Mean regret (lower is better)", x + 6.2, y + h + 0.25, 6.0, 0.35, 8, SLATE, align=PP_ALIGN.CENTER)
    textbox(slide, "Success rate", x + 0.2, y - 0.45, 4.0, 0.35, 8, SLATE)
    bullets(slide, ["The oracle is an evaluation upper bound, not an input to the selector.", "V5 methods reduce regret compared with simple heuristics.", "Shortest-tree is weak: smaller BTs are not necessarily more executable."], 21.4, 4.2, 10.0, 5.4, 16)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Selector Robustness Across Benchmarks", "Mean regret heatmap")
    cols, mat = selector_matrix("regret")
    heatmap(slide, cols, mat, 1.4, 3.0, 22.5, 10.2, "Mean regret by selector and benchmark", lower_is_better=True)
    bullets(slide, ["Heatmap highlights where each selector is reliable or brittle.", "Hard cases make physical reliability more important.", "The final discussion should reference every figure included in the thesis text."], 24.8, 4.2, 7.0, 5.5, 15)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "V5 Ablation Study", "Which component matters?")
    ab_data = []
    for r in ablations[:12]:
        regret = n(first(r, ["mean_regret", "regret"]))
        if regret is not None:
            ab_data.append((first(r, ["selector"]), regret))
    bar_chart(slide, ab_data, 1.4, 3.1, 17.5, 10.2, "Mean regret by ablation variant", PURPLE, lower_is_better=True)
    card(slide, "Finding 1", "Removing symbolic reliability weakens rejection of invalid or incomplete BTs.", 20.2, 3.5, 10.2, 2.4, GREEN)
    card(slide, "Finding 2", "Removing simulation reliability makes physically risky strategies harder to distinguish.", 20.2, 6.4, 10.2, 2.4, AMBER)
    card(slide, "Finding 3", "V5 works because signals are complementary, not because of a single module.", 20.2, 9.3, 10.2, 2.4, BLUE)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Simulation Validation", "Why Isaac Gym is used")
    outcome_bars = []
    for r in outcomes:
        total = (n(first(r, ["sim_success"])) or 0) + (n(first(r, ["sim_failure"])) or 0)
        outcome_bars.append((first(r, ["benchmark"]), total))
    bar_chart(slide, outcome_bars, 1.4, 3.3, 14.0, 8.5, "Simulated candidate outcomes", CYAN)
    bullets(slide, ["KIOS checks abstract symbolic execution.", "Isaac Gym exposes physical issues: unstable support, edge placement, displacement and contact violations.", "Simulation metrics provide supervision for learned selectors and evidence for evaluation."], 17.8, 4.0, 12.8, 5.3, 17)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "GUI Demonstration", "End-to-end visual workflow")
    step(slide, "Task definition", 1.8, 4.0, 5.0, 1.2)
    step(slide, "Multiple BT candidates", 8.0, 4.0, 5.8, 1.2)
    step(slide, "Selector score table", 15.2, 4.0, 5.8, 1.2)
    step(slide, "Selected / Oracle BT", 22.4, 4.0, 5.8, 1.2, RGBColor(220, 252, 231))
    arrow(slide, 6.8, 4.6, 8.0, 4.6); arrow(slide, 13.8, 4.6, 15.2, 4.6); arrow(slide, 21.0, 4.6, 22.4, 4.6)
    card(slide, "Web demo", "Shows task, candidate BTs, selector scores, top-down placement view, and replay commands.", 3.0, 8.0, 12.0, 3.0, BLUE)
    card(slide, "Visual replay", "Provides a presentation-friendly Isaac Gym style execution view.", 17.0, 8.0, 11.0, 3.0, AMBER)
    textbox(slide, "Demo URL: http://localhost:8091/", 6.0, 12.5, 18.0, 0.5, 16, NAVY, True, PP_ALIGN.CENTER)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Project Management", "How complexity was handled")
    table(slide, ["Challenge", "Decision / Management Response"], [
        ["Many moving parts", "Separated pipeline into candidate generation, KIOS evaluation, selector training, simulation and visualization"],
        ["Expensive simulation", "Ran benchmark stages incrementally and reused cached results"],
        ["Changing research direction", "Used staged versions V1, V2, V2B, V4B, V5 and V6 to keep progress measurable"],
        ["API instability", "Kept GPT single-vs-multi as a pilot and focused main claims on completed benchmark results"],
        ["Writing and reporting", "Prepared final figures, method chapter, experiment chapter and defense outline before further code changes"],
    ], 1.2, 3.0, 30.8, 9.5)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Main Contributions")
    card(slide, "System", "An end-to-end BT selection pipeline combining KIOS, Transformer encoding, V5 selection and Isaac Gym validation.", 1.8, 3.2, 9.2, 4.0, BLUE)
    card(slide, "Benchmarks", "Four task groups covering supported-by relations, stacking probes, physical strategy differences and hard cases.", 12.3, 3.2, 9.2, 4.0, GREEN)
    card(slide, "Algorithm", "A hybrid learned V5 selector that fuses selector agreement, symbolic reliability and physical reliability.", 22.8, 3.2, 9.2, 4.0, PURPLE)
    card(slide, "Evaluation", "Final tables, vector-based figures, ablation analysis and simulation-backed validation for thesis reporting.", 7.0, 9.0, 19.8, 3.2, AMBER)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs)); title(slide, "Future Work")
    bullets(slide, ["Complete a larger GPT single-vs-multi comparison once API limits are no longer a bottleneck.", "Add more long-horizon manipulation tasks and randomized initial states.", "Explore graph or hypergraph encoders for richer BT and object-relation structure.", "Close the loop with BT repair and online feedback from failed executions.", "Transfer the selected BTs to more realistic robot settings."], 2.0, 3.0, 27.0, 7.5, 19)
    footer(slide, page); page += 1

    slide = prs.slides.add_slide(blank(prs))
    textbox(slide, "Thank You", 1.4, 5.7, 31.0, 1.0, 36, NAVY, True, PP_ALIGN.CENTER)
    textbox(slide, "Questions?", 1.4, 7.2, 31.0, 0.8, 24, BLUE, True, PP_ALIGN.CENTER)
    textbox(slide, "KIOS + Transformer BT Selector for Robot Task Planning", 1.4, 11.8, 31.0, 0.5, 14, SLATE, align=PP_ALIGN.CENTER)
    footer(slide, page)

    return prs


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--author", default="Your Name")
    parser.add_argument("--department", default="Your Department")
    parser.add_argument("--date", default="2026")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    prs = build_deck(args)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print({"result": "success", "output": str(args.output), "slides": len(prs.slides)})


if __name__ == "__main__":
    main()
